"""Relatórios fechados (semana / mês) para o gestor, um arquivo por período."""
import os
from collections import defaultdict
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from export_xlsx import BOLD, HEAD, HOURS, LINE, MESES, MUTED, SUBTLE, WHITE_BOLD, _dur, _header, _widths, _work_dates

DIAS = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta"]
KIND = {"agendada": "agendada", "chat": "link do chat", "avulsa": "fora da agenda", "presencial": "presencial"}
GREEN, RED = Font(color="006300"), Font(color="B42318")


def week_bounds(d):
    a = d - timedelta(days=d.weekday())
    return a, a + timedelta(days=4)


def month_bounds(d):
    a = d.replace(day=1)
    b = date(a.year + (a.month == 12), a.month % 12 + 1, 1) - timedelta(days=1)
    return a, b


def report_name(kind, start, end):
    if kind == "semanal":
        return f"Semanal {start:%Y-%m-%d} ({start:%d-%m} a {end:%d-%m}).xlsx"
    return f"Mensal {start:%Y-%m} ({MESES[start.month - 1]} {start.year}).xlsx"


class Ctx:
    def __init__(self, data):
        self.rows, self.cats, self.H = data["rows"], data["categories"], data["hoursPerWeek"]
        self.off = set(data["daysOff"])
        self.first, self.today = date.fromisoformat(data["firstDay"]), date.fromisoformat(data["today"])
        self.data = data
        for r in self.rows:
            r["d"] = date.fromisoformat(r["date"])
        self.missed = data.get("missed", [])
        for m in self.missed:
            m["d"] = date.fromisoformat(m["date"])

    def period(self, a, b):
        rows = [r for r in self.rows if a <= r["d"] <= b]
        days = _work_dates(a, b, self.first, self.today, self.off)
        cap = len(days) * self.H * 12
        tot = sum(r["min"] for r in rows)
        missed = [m for m in self.missed if a <= m["d"] <= b]
        return {
            "missed": missed, "missed_min": sum(m.get("min", 0) for m in missed),
            "rows": rows, "days": days, "cap": cap, "total": tot,
            "rec": sum(r["min"] for r in rows if r["rec"]),
            "off": sum(r["min"] for r in rows if r["kind"] in ("chat", "avulsa")),
            "n_off": sum(1 for r in rows if r["kind"] in ("chat", "avulsa")),
            "cat": {c: sum(r["min"] for r in rows if r["cat"] == c) for c in self.cats},
        }


def _pct(m, cap):
    return m / cap if cap else 0


def _delta_cell(ws, row, col, now, before, lower_is_good=True):
    """Diferença em pontos percentuais da capacidade, verde quando melhora."""
    if before is None:
        ws.cell(row, col, "–")
        return
    d = now - before
    c = ws.cell(row, col, d)
    c.number_format = '+0%;-0%;0%'
    if round(d, 2):
        c.font = GREEN if (d < 0) == lower_is_good else RED


def build_report(data, kind, start, end, path):
    ctx = Ctx(data)
    cur = ctx.period(start, end)
    # base de comparação: 4 semanas anteriores (semanal) ou mês anterior (mensal)
    if kind == "semanal":
        base = ctx.period(start - timedelta(weeks=4), start - timedelta(days=3))
        base_label = "média 4 semanas anteriores"
        title = f"Semana de {start:%d/%m} a {end:%d/%m/%Y}"
    else:
        pa, pb = month_bounds(start - timedelta(days=1))
        base = ctx.period(pa, pb)
        base_label = f"{MESES[pa.month - 1]}"
        title = f"{MESES[start.month - 1].capitalize()} de {start.year}"
    has_base = bool(base["days"])
    cap = cur["cap"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = f"Gestão do Tempo — {title}"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = (f"{data.get('owner') or 'Relatório'} · {len(cur['days'])} dias úteis · capacidade {ctx.H}h/semana "
                f"({cap / 60:.1f}h no período) · medido pelo Meet, Google Agenda e microfone · gerado {data['generated']}")
    ws["A2"].font = MUTED

    # Indicadores
    _header(ws, 4, ["Indicador", "Horas", "% do período", f"vs {base_label}"])
    kpis = [
        ("Em reuniões", cur["total"], base["total"]),
        ("Recorrentes (pré-travado)", cur["rec"], base["rec"]),
        ("Fora da agenda (chat / link direto)", cur["off"], base["off"]),
        ("Agendadas em que não entrei", cur["missed_min"], base["missed_min"]),
        ("Tempo livre para foco e demandas", max(0, cap - cur["total"]), max(0, base["cap"] - base["total"])),
    ]
    for i, (label, m, bm) in enumerate(kpis, 5):
        ws.cell(i, 1, label).font = BOLD if i == 5 else Font()
        ws.cell(i, 2, _dur(m)).number_format = HOURS
        ws.cell(i, 3, _pct(m, cap)).number_format = "0%"
        bpct = _pct(bm, base["cap"]) if has_base else None
        _delta_cell(ws, i, 4, _pct(m, cap), bpct, lower_is_good=not label.startswith("Tempo livre"))
        for col in range(1, 5):
            ws.cell(i, col).border = LINE
    ws.cell(10, 1, f"{cur['n_off']} calls fora da agenda · {len(cur['missed'])} reuniões aceitas em que não entrei "
                   f"(não somam em 'Em reuniões'; detalhe na aba 'Sem participação')").font = MUTED
    ws.cell(11, 1, "Comparação em pontos percentuais da capacidade (verde = melhor).").font = MUTED

    # Compilado por tema
    r0 = 13
    _header(ws, r0, ["Compilado por tema", "Horas", "% do período", f"vs {base_label}", "Recorrente"])
    row = r0 + 1
    for c in sorted(ctx.cats, key=lambda c: -cur["cat"][c]):
        if not cur["cat"][c] and not base["cat"][c]:
            continue
        ws.cell(row, 1, c)
        ws.cell(row, 2, _dur(cur["cat"][c])).number_format = HOURS
        ws.cell(row, 3, _pct(cur["cat"][c], cap)).number_format = "0%"
        _delta_cell(ws, row, 4, _pct(cur["cat"][c], cap), _pct(base["cat"][c], base["cap"]) if has_base else None)
        ws.cell(row, 5, _dur(sum(r["min"] for r in cur["rows"] if r["cat"] == c and r["rec"]))).number_format = HOURS
        for col in range(1, 6):
            ws.cell(row, col).border = LINE
        row += 1
    for label, m in (("Tempo focado / livre", max(0, cap - cur["total"])), ("Total (capacidade)", cap)):
        ws.cell(row, 1, label)
        ws.cell(row, 2, _dur(m)).number_format = HOURS
        ws.cell(row, 3, _pct(m, cap)).number_format = "0%"
        for col in range(1, 6):
            ws.cell(row, col).font = BOLD
            ws.cell(row, col).fill = SUBTLE
        row += 1

    # Distribuição: por dia (semanal) ou por semana (mensal)
    row += 1
    if kind == "semanal":
        _header(ws, row, ["Por dia", "Em reuniões", "% do dia"])
        row += 1
        day_cap = ctx.H * 12
        for i, nome in enumerate(DIAS):
            d = start + timedelta(days=i)
            m = sum(r["min"] for r in cur["rows"] if r["d"] == d)
            ws.cell(row, 1, f"{nome} {d:%d/%m}" + (" (sem expediente)" if d.isoformat() in ctx.off else ""))
            ws.cell(row, 2, _dur(m)).number_format = HOURS
            ws.cell(row, 3, _pct(m, day_cap)).number_format = "0%"
            for col in range(1, 4):
                ws.cell(row, col).border = LINE
            row += 1
    else:
        _header(ws, row, ["Por semana", "Em reuniões", "% da semana", "Fora da agenda"])
        row += 1
        chart_first = row
        wk = start - timedelta(days=start.weekday())
        while wk <= end:
            a, b = max(wk, start), min(wk + timedelta(days=4), end)
            wk += timedelta(weeks=1)
            if a > b:  # semana que só encosta no mês pelo fim de semana
                continue
            p = ctx.period(a, b)
            ws.cell(row, 1, f"{a:%d/%m} – {b:%d/%m}")
            ws.cell(row, 2, _dur(p["total"])).number_format = HOURS
            ws.cell(row, 3, _pct(p["total"], p["cap"])).number_format = "0%"
            ws.cell(row, 4, _dur(p["off"])).number_format = HOURS
            for col in range(1, 5):
                ws.cell(row, col).border = LINE
            row += 1
        ch = BarChart()
        ch.title, ch.legend = "Horas em reunião por semana", None
        ch.y_axis.number_format = '[h]"h"'
        ch.add_data(Reference(ws, min_col=2, min_row=chart_first, max_row=row - 1))
        ch.set_categories(Reference(ws, min_col=1, min_row=chart_first, max_row=row - 1))
        ch.height, ch.width = 7, 14
        ws.add_chart(ch, "G12")
    _widths(ws, [38, 13, 13, 22, 12])

    # Atividades (formato da planilha do gestor)
    s = wb.create_sheet("Atividades")
    _header(s, 1, ["Categoria", "Atividade", "Horas", "Vezes", "Tipo"])
    acts = defaultdict(lambda: {"m": 0, "n": 0, "rec": False, "kinds": set()})
    for r in cur["rows"]:
        a = acts[(r["cat"], r["title"])]
        a["m"] += r["min"]; a["n"] += 1; a["rec"] |= r["rec"]; a["kinds"].add(r["kind"])
    row = 2
    for c in sorted(ctx.cats, key=lambda c: -cur["cat"][c]):
        items = sorted(((t, a) for (cc, t), a in acts.items() if cc == c), key=lambda x: -x[1]["m"])
        for t, a in items:
            tipo = ["recorrente" if a["rec"] else "pontual"] + [KIND[k] for k in a["kinds"] if k != "agendada"]
            for col, v in enumerate([c, t, _dur(a["m"]), a["n"], ", ".join(tipo)], 1):
                s.cell(row, col, v).border = LINE
            s.cell(row, 3).number_format = HOURS
            row += 1
        if items:
            row += 1
    _widths(s, [22, 56, 10, 8, 30])
    s.freeze_panes = "A2"

    # Calls fora da agenda
    s = wb.create_sheet("Fora da agenda")
    _header(s, 1, ["Data", "Início", "Duração", "Call", "Origem"])
    off_rows = sorted((r for r in cur["rows"] if r["kind"] in ("chat", "avulsa")), key=lambda r: (r["date"], r["start"]))
    for i, r in enumerate(off_rows, 2):
        for col, v in enumerate([r["d"], r["start"], _dur(r["min"]), r["title"], KIND[r["kind"]]], 1):
            s.cell(i, col, v).border = LINE
        s.cell(i, 1).number_format = "dd/mm/yyyy"
        s.cell(i, 3).number_format = HOURS
    if not off_rows:
        s.cell(2, 1, "Nenhuma call fora da agenda no período.").font = MUTED
    _widths(s, [12, 8, 10, 52, 16])

    # Agendadas sem participação: aceitas, com outras pessoas, sem entrada no Meet
    s = wb.create_sheet("Sem participação")
    _header(s, 1, ["Data", "Início", "Duração", "Reunião aceita em que não entrei", "Tema", "Recorrente"])
    miss = sorted(cur["missed"], key=lambda m: (m["date"], m["start"]))
    for i, m in enumerate(miss, 2):
        vals = [m["d"], m["start"], _dur(m.get("min", 0)), m["title"], m.get("cat", ""), "sim" if m.get("rec") else "não"]
        for col, v in enumerate(vals, 1):
            s.cell(i, col, v).border = LINE
        s.cell(i, 1).number_format = "dd/mm/yyyy"
        s.cell(i, 3).number_format = HOURS
    last = len(miss) + 1
    if not miss:
        s.cell(2, 1, "Participei de todas as reuniões aceitas no período.").font = MUTED
        last = 2
    s.cell(last + 2, 1, "Calls pelo celular, outro navegador ou perfil do Chrome não são detectadas e podem aparecer aqui.").font = MUTED
    # quais reuniões mais se repetem: ajuda a decidir o que recusar ou sair da recorrência
    by = defaultdict(lambda: [0, 0.0])
    for m in miss:
        by[m["title"]][0] += 1; by[m["title"]][1] += m.get("min", 0)
    if by:
        for i, v in enumerate(["Reunião", "Vezes", "Horas"], 8):
            c = s.cell(1, i, v); c.fill, c.font = HEAD, WHITE_BOLD
        for i, (t, (n, mm)) in enumerate(sorted(by.items(), key=lambda kv: (-kv[1][0], -kv[1][1])), 2):
            s.cell(i, 8, t).border = LINE
            s.cell(i, 9, n).border = LINE
            c = s.cell(i, 10, _dur(mm)); c.number_format = HOURS; c.border = LINE
    _widths(s, [12, 8, 10, 50, 20, 11, 3, 44, 8, 10])
    s.freeze_panes = "A2"

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path


def due_reports(today):
    """Períodos já encerrados: a última semana (seg–sex) e o último mês completos, mais o anterior de cada,
    para recuperar caso o PC estivesse desligado na segunda."""
    out = []
    last_fri = today - timedelta(days=(today.weekday() - 4) % 7 or 7)  # sexta mais recente antes de hoje
    for k in (0, 1):
        a, b = week_bounds(last_fri - timedelta(weeks=k))
        out.append(("semanal", a, b))
    ma, _ = month_bounds(today)
    for k in (1, 2):
        pa, pb = month_bounds(ma - timedelta(days=1))
        out.append(("mensal", pa, pb))
        ma = pa
    return out
