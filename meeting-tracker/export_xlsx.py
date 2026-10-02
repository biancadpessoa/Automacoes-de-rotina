"""Planilha para compartilhar com o gestor, no formato Categoria > Atividade > Horas/semana."""
import os
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="1F3A5F")
SUBTLE = PatternFill("solid", fgColor="F0EFEC")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
MUTED = Font(color="6B6A64", italic=True)
LINE = Border(bottom=Side(style="thin", color="D9D8D2"))
HOURS = '[h]"h"mm'  # duração em fração de dia
MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]


def _dur(minutes):
    return minutes / 1440  # o Excel guarda duração como fração de dia


def _work_dates(start, end, first, today, off):
    d, out = start, []
    while d <= end:
        if d.weekday() < 5 and first <= d <= today and d.isoformat() not in off:
            out.append(d)
        d += timedelta(days=1)
    return out


def _header(ws, row, values):
    for i, v in enumerate(values, 1):
        c = ws.cell(row, i, v)
        c.fill, c.font, c.alignment = HEAD, WHITE_BOLD, Alignment(vertical="center", wrap_text=True)


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(data, path, weeks_back=4):
    rows, cats, H = data["rows"], data["categories"], data["hoursPerWeek"]
    off = set(data["daysOff"])
    first, today = date.fromisoformat(data["firstDay"]), date.fromisoformat(data["today"])
    cap_day = H * 60 / 5
    for r in rows:
        r["d"] = date.fromisoformat(r["date"])

    # Período do resumo: últimas N semanas completas (a atual ainda está em andamento)
    this_monday = today - timedelta(days=today.weekday())
    p_start = this_monday - timedelta(weeks=weeks_back)
    p_end = this_monday - timedelta(days=1)
    p_rows = [r for r in rows if p_start <= r["d"] <= p_end]
    p_weeks = len(_work_dates(p_start, p_end, first, today, off)) / 5 or 1  # semanas "cheias" equivalentes
    per_week = lambda m: m / p_weeks

    wb = Workbook()

    # ---------- Resumo (formato da planilha do gestor) ----------
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "Gestão do Tempo"
    ws["A1"].font = Font(bold=True, size=16)
    ws["B1"] = data.get("owner", "")
    ws["B1"].font = Font(size=12)
    ws["A2"] = (f"Média por semana de {p_start:%d/%m} a {p_end:%d/%m/%Y} ({weeks_back} semanas completas) · "
                f"medido automaticamente pelo Meet, Google Agenda e microfone · atualizado {data['generated']}")
    ws["A2"].font = MUTED

    _header(ws, 4, ["Categoria", "Atividade", "Horas/semana", "Vezes", "Tipo"])
    acts = defaultdict(lambda: {"m": 0, "n": 0, "rec": False, "kinds": set()})
    for r in p_rows:
        a = acts[(r["cat"], r["title"])]
        a["m"] += r["min"]; a["n"] += 1; a["rec"] |= r["rec"]; a["kinds"].add(r["kind"])
    cat_total = defaultdict(float)
    for (c, _), a in acts.items():
        cat_total[c] += a["m"]
    row = 5
    kind_label = {"chat": "link do chat", "avulsa": "fora da agenda", "presencial": "presencial"}
    for c in sorted(cat_total, key=lambda c: -cat_total[c]):
        items = sorted(((t, a) for (cc, t), a in acts.items() if cc == c), key=lambda x: -x[1]["m"])
        for t, a in items:
            ws.cell(row, 1, c)
            ws.cell(row, 2, t)
            ws.cell(row, 3, _dur(per_week(a["m"]))).number_format = HOURS
            ws.cell(row, 4, a["n"])
            tipo = ["recorrente"] if a["rec"] else ["pontual"]
            tipo += [kind_label[k] for k in a["kinds"] if k in kind_label]
            ws.cell(row, 5, ", ".join(tipo))
            for col in range(1, 6):
                ws.cell(row, col).border = LINE
            row += 1
        row += 1  # linha em branco entre categorias, como no original

    # Compilado por tema, ao lado (colunas G–J)
    for i, v in enumerate(["Compilado por tema", "Horas/semana", f"% das {H}h", "Recorrente"], 7):
        c = ws.cell(4, i, v)
        c.fill, c.font, c.alignment = HEAD, WHITE_BOLD, Alignment(vertical="center", wrap_text=True)
    total_m = sum(r["min"] for r in p_rows)
    rec_m = sum(r["min"] for r in p_rows if r["rec"])
    off_m = sum(r["min"] for r in p_rows if r["kind"] in ("chat", "avulsa"))
    cap = H * 60
    focus = max(0, cap - per_week(total_m))
    r2 = 5
    for c in sorted(cat_total, key=lambda c: -cat_total[c]):
        m = per_week(cat_total[c])
        rec = per_week(sum(r["min"] for r in p_rows if r["cat"] == c and r["rec"]))
        ws.cell(r2, 7, c)
        ws.cell(r2, 8, _dur(m)).number_format = HOURS
        ws.cell(r2, 9, m / cap).number_format = "0%"
        ws.cell(r2, 10, _dur(rec)).number_format = HOURS
        for col in range(7, 11):
            ws.cell(r2, col).border = LINE
        r2 += 1
    for label, m, fill in (("Tempo focado / livre", focus, SUBTLE), ("Total", cap, None)):
        ws.cell(r2, 7, label).font = BOLD
        ws.cell(r2, 8, _dur(m)).number_format = HOURS
        ws.cell(r2, 9, m / cap).number_format = "0%"
        for col in range(7, 11):
            ws.cell(r2, col).font = BOLD
            if fill:
                ws.cell(r2, col).fill = fill
        r2 += 1
    r2 += 1
    notes = [
        f"{per_week(rec_m) / cap:.0%} da semana pré-travado em reuniões recorrentes",
        f"{per_week(total_m) / cap:.0%} da semana em reuniões no total",
        f"{per_week(off_m) / 60:.1f}h/semana em calls fora da agenda (links do chat ou diretos)",
        f"{focus / cap:.0%} livre para foco e demandas",
    ]
    for n in notes:
        ws.cell(r2, 7, "• " + n)
        r2 += 1
    _widths(ws, [22, 52, 13, 8, 26, 3, 26, 13, 11, 12])
    ws.freeze_panes = "A5"

    # ---------- Semanal e Mensal ----------
    def period_sheet(title, key, label, bounds):
        s = wb.create_sheet(title)
        periods = sorted({key(r) for r in rows})
        _header(s, 1, [label, "Dias úteis", "Capacidade", "Em reuniões", "% reuniões", "Recorrente",
                       "Fora da agenda", "Tempo livre"] + cats)
        for i, p in enumerate(periods, 2):
            pr = [r for r in rows if key(r) == p]
            a, b = bounds(p)
            days = len(_work_dates(a, b, first, today, off))
            capm = days * cap_day
            tot = sum(r["min"] for r in pr)
            vals = [f"{MESES[a.month - 1]} {a.year}" if title == "Mensal" else f"{a:%d/%m} – {b - timedelta(days=2):%d/%m/%Y}", days,
                    _dur(capm), _dur(tot), tot / capm if capm else 0,
                    _dur(sum(r["min"] for r in pr if r["rec"])),
                    _dur(sum(r["min"] for r in pr if r["kind"] in ("chat", "avulsa"))),
                    _dur(max(0, capm - tot))] + [_dur(sum(r["min"] for r in pr if r["cat"] == c)) for c in cats]
            for j, v in enumerate(vals, 1):
                c = s.cell(i, j, v)
                c.border = LINE
                if j in (3, 4, 6, 7, 8) or j > 8:
                    c.number_format = HOURS
                elif j == 5:
                    c.number_format = "0%"
        _widths(s, [22, 10, 12, 12, 11, 12, 13, 12] + [14] * len(cats))
        s.freeze_panes = "B2"
        # gráfico empilhado por tema
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
        ch.title, ch.y_axis.title = f"Horas em reunião por {label.lower()}", "horas"
        ch.y_axis.number_format = '[h]"h"'
        ch.add_data(Reference(s, min_col=9, max_col=8 + len(cats), min_row=1, max_row=len(periods) + 1), titles_from_data=True)
        ch.set_categories(Reference(s, min_col=1, min_row=2, max_row=len(periods) + 1))
        ch.height, ch.width = 9, 22
        s.add_chart(ch, f"A{len(periods) + 4}")
        return s

    def week_bounds(w):
        y, n = map(int, w.split("-W"))
        a = date.fromisocalendar(y, n, 1)
        return a, a + timedelta(days=6)

    def month_bounds(m):
        y, n = map(int, m.split("-"))
        a = date(y, n, 1)
        b = (date(y + (n == 12), n % 12 + 1, 1)) - timedelta(days=1)
        return a, b

    period_sheet("Semanal", lambda r: r["week"], "Semana", week_bounds)
    period_sheet("Mensal", lambda r: r["date"][:7], "Mês", month_bounds)

    # ---------- Registro de calls ----------
    s = wb.create_sheet("Calls")
    _header(s, 1, ["Data", "Início", "Fim", "Duração", "Atividade", "Tema", "Tipo", "Recorrente"])
    tipo = {"agendada": "agendada", "chat": "link do chat", "avulsa": "fora da agenda", "presencial": "presencial"}
    for i, r in enumerate(sorted(rows, key=lambda r: (r["date"], r["start"]), reverse=True), 2):
        vals = [r["d"], r["start"], r["end"], _dur(r["min"]), r["title"], r["cat"], tipo[r["kind"]], "sim" if r["rec"] else "não"]
        for j, v in enumerate(vals, 1):
            s.cell(i, j, v)
        s.cell(i, 1).number_format = "dd/mm/yyyy"
        s.cell(i, 4).number_format = HOURS
    s.auto_filter.ref = f"A1:H{len(rows) + 1}"
    s.freeze_panes = "A2"
    _widths(s, [12, 8, 8, 10, 52, 22, 16, 11])

    # ---------- Agendadas sem participação ----------
    if data["calendar"].startswith("agenda conectada"):
        s = wb.create_sheet("Sem participação")
        _header(s, 1, ["Data", "Início", "Reunião aceita em que não entrei"])
        for i, m in enumerate(sorted(data["missed"], key=lambda m: (m["date"], m["start"]), reverse=True), 2):
            s.cell(i, 1, date.fromisoformat(m["date"])).number_format = "dd/mm/yyyy"
            s.cell(i, 2, m["start"])
            s.cell(i, 3, m["title"])
        s.cell(len(data["missed"]) + 3, 1, "Calls pelo celular ou outro perfil do navegador não são detectadas.").font = MUTED
        _widths(s, [12, 8, 60])

    wb.save(path)
    return path
